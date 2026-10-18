"""데모·테스트용 샘플 브리핑 자료를 만든다 (일부러 문제점을 포함).

    python samples/make_samples.py [출력폴더]
"""
from __future__ import annotations

import sys
from pathlib import Path

import pymupdf

SLIDES = [
    {
        "title": "2025 스마트시티 데이터 거버넌스 고도화 추진계획",
        "body": [
            "■ 추진배경: 디지털 트윈 기반 SoS(System of Systems) 아키텍처 내재화를 통한",
            "   데이터 레이크하우스 interoperability 확보 및 KPI 정합성 제고",
            "■ 기대효과: TCO 23.7% 절감, ROI 312%, MAU 1.8M 달성(YoY +41.2%)",
        ],
        "footnote": "※ 상세 산출근거는 별첨 7-3 참조 (본 자료 미포함)",
    },
    {
        "title": "추진 현황 및 향후 일정",
        "body": [
            "1단계(’24.Q3~’25.Q1): PoC 완료, MVP 고도화 및 ISMS-P 인증 대응, 레거시 DW 마이그레이션 착수,",
            "API 게이트웨이 표준화, 메타데이터 카탈로그 구축, 데이터 품질관리(DQM) 체계 수립, 개인정보 가명처리",
            "가이드라인 개정, 클라우드 네이티브 전환(K8s), MSA 리팩토링, CI/CD 파이프라인 구축, 관제센터 연계,",
            "2단계(’25.Q2~): 전면 확산 — 시민 체감형 서비스 발굴(미정)",
        ],
        "footnote": "문의: 데이터정책과 (내선 미기재)",
    },
]


# 일반 시민 관점의 지적을 일부 반영한 수정본 (수정본 비교 데모용)
SLIDES_V2 = [
    {
        "title": "우리 동네 스마트시티, 무엇이 좋아지나요?",
        "body": [
            "■ 버스 도착, 주차 빈자리, 미세먼지 정보를 한 앱에서 확인",
            "■ 운영비(TCO) 23.7% 절감, ROI 312% → 생활 서비스에 다시 투자",
            "■ 데이터 관리 체계(거버넌스) 정비로 개인정보 보호 강화",
        ],
        "footnote": "※ 비용 산출 근거: 시 홈페이지 > 정책자료실 (문의 02-123-4567)",
    },
    {
        "title": "언제부터 쓸 수 있나요?",
        "body": [
            "2025년 3월: 시범 운영 (3개 동)",
            "2025년 9월: 전체 동 확대",
            "신청: 앱 설치 또는 주민센터 방문 (세부 절차는 추후 공지)",
        ],
        "footnote": "문의: 데이터정책과 02-123-4567",
    },
]


def make_pdf(path: Path, slides: list[dict] | None = None, footnote_size: float = 8) -> Path:
    doc = pymupdf.open()
    for slide in slides or SLIDES:
        page = doc.new_page(width=960, height=540)
        page.draw_rect(pymupdf.Rect(0, 0, 960, 90), color=None, fill=(0.12, 0.2, 0.45))
        page.insert_text((40, 58), slide["title"], fontname="korea", fontsize=28, color=(1, 1, 1))
        y = 150
        for line in slide["body"]:
            page.insert_text((40, y), line, fontname="korea", fontsize=17)
            y += 34
        # 일부러 너무 작고 흐린 각주
        page.insert_text((40, 515), slide["footnote"], fontname="korea", fontsize=footnote_size, color=(0.7, 0.7, 0.7))
    doc.save(path)
    return path


def make_pptx(path: Path) -> Path | None:
    try:
        from pptx import Presentation
        from pptx.util import Inches, Pt
    except ImportError:
        return None
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    for slide in SLIDES:
        s = prs.slides.add_slide(prs.slide_layouts[5])
        s.shapes.title.text = slide["title"]
        box = s.shapes.add_textbox(Inches(0.6), Inches(1.8), Inches(12), Inches(4)).text_frame
        box.word_wrap = True
        for i, line in enumerate(slide["body"]):
            para = box.paragraphs[0] if i == 0 else box.add_paragraph()
            para.text = line
            para.font.size = Pt(18)
        foot = s.shapes.add_textbox(Inches(0.6), Inches(6.9), Inches(12), Inches(0.4)).text_frame
        foot.text = slide["footnote"]
        foot.paragraphs[0].font.size = Pt(8)
    prs.save(path)
    return path


HTML = """<!doctype html><html lang="ko"><head><meta charset="utf-8">
<style>body{font-family:sans-serif;margin:0;color:#222}
header{background:#1e3370;color:#fff;padding:32px 48px}
section{padding:24px 48px} .tiny{font-size:9px;color:#bbb}
table{border-collapse:collapse} td,th{border:1px solid #999;padding:4px 8px;font-size:12px}</style>
</head><body>
<header><h1>공공 마이데이터 연계 플랫폼 대국민 안내</h1></header>
<section><h2>서비스 개요</h2>
<p>본 서비스는 OAuth 2.0 기반 전송요구권 행사를 통해 정보주체의 개인신용정보를 API 방식으로 제3자에게 전송하는 체계입니다.</p>
<table><tr><th>구분</th><th>1단계</th><th>2단계</th><th>3단계</th><th>4단계</th><th>5단계</th></tr>
<tr><td>대상기관</td><td>12</td><td>37</td><td>81</td><td>140</td><td>TBD</td></tr>
<tr><td>연계항목</td><td>214</td><td>633</td><td>1,207</td><td>2,950</td><td>TBD</td></tr></table>
<p class="tiny">※ 개인정보 제3자 제공 동의 철회는 별도 절차에 따름</p>
</section>
<section><h2>신청 방법</h2><p>추후 공지 예정</p></section>
</body></html>
"""


def make_html(path: Path) -> Path:
    path.write_text(HTML, encoding="utf-8")
    return path


def make_all(out: Path) -> list[Path]:
    out.mkdir(parents=True, exist_ok=True)
    made = [make_pdf(out / "sample_briefing.pdf"), make_html(out / "sample_notice.html"),
            make_pdf(out / "sample_briefing_v2.pdf", SLIDES_V2, footnote_size=12)]
    pptx = make_pptx(out / "sample_briefing.pptx")
    if pptx:
        made.append(pptx)
    return made


if __name__ == "__main__":
    target = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(__file__).resolve().parent
    for p in make_all(target):
        print(p)
